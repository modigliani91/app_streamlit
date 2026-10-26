"""Téléchargement du fichier exact, conditionné par reçu de publication et SHA-256."""
from dataclasses import dataclass, field, asdict
from pathlib import Path
import hashlib
import json
from exceptions import ReportUnavailable
from file_manager import RunLock, atomic_json, sha256_file
from progress import utc_now

@dataclass
class PipelineResult:
    success: bool
    run_id: str
    report_path: str | None = None
    customers_discovered: int = 0
    customers_processed: int = 0
    customers_failed: int = 0
    documents_processed: int = 0
    conform_count: int = 0
    non_conform_count: int = 0
    recheck_count: int = 0
    reference_not_found_count: int = 0
    manual_review_count: int = 0
    errors: list[str] = field(default_factory=list)
    elapsed_seconds: float = 0.0
    qwen_calls: int = 0
    pipeline_version: str = ""
    matching_version: str = ""
    report_sha256: str = ""
    warnings: bool = False

    def to_dict(self): return asdict(self)

def invalidate_report(settings, run_id):
    # Remove the receipt first; a failed run cannot advertise an old report.
    atomic_json(settings.report_receipt, {"run_id":run_id,"status":"PENDING","updated_at":utc_now()})
    if settings.report_path.exists():
        archive=settings.output_dir/"report_history"/run_id/settings.report_filename
        archive.parent.mkdir(parents=True,exist_ok=True)
        settings.report_path.replace(archive)

def validate_report(path, started_ns, customer_count):
    from openpyxl import load_workbook
    path=Path(path)
    if not path.is_file() or path.stat().st_size<=0:
        raise ReportUnavailable("Le traitement s'est terminé mais le rapport Excel n'a pas pu être généré.")
    # Absence at start + atomic creation + hash are primary freshness proofs.
    # Allow 2 seconds for filesystems with coarse timestamps.
    if path.stat().st_mtime_ns < started_ns-2_000_000_000:
        raise ReportUnavailable("La date du rapport ne correspond pas au traitement actuel.")
    workbook=load_workbook(path,read_only=True,data_only=False)
    try:
        required={"SUMMARY","CUSTOMER_CONTROLS","DOCUMENT_CONTROLS","FATCA_DETAILS","SIGNATURE_DETAILS","DATABASE_MATCHING"}
        if not required.issubset(workbook.sheetnames):
            raise ReportUnavailable("Le rapport Excel est incomplet.")
        if workbook["CUSTOMER_CONTROLS"].max_row-1 != customer_count:
            raise ReportUnavailable("Le nombre de clients du rapport ne correspond pas à l'inventaire.")
    finally: workbook.close()
    return sha256_file(path)

def publish_receipt(settings, result, input_sha256):
    atomic_json(settings.report_receipt,{"run_id":result.run_id,"status":"COMPLETED",
        "input_sha256":input_sha256,"report_path":str(settings.report_path),
        "report_sha256":result.report_sha256,"updated_at":utc_now()})

def download_report(settings, run_id, input_sha256):
    # Lock spans receipt validation and byte reading; a new job cannot replace it mid-read.
    with RunLock(settings.lock_path, "download-"+run_id):
        try:
            receipt=json.loads(settings.report_receipt.read_text())
            metadata=json.loads((settings.metadata_dir/(run_id+".json")).read_text())
            if (receipt.get("status")!="COMPLETED" or receipt.get("run_id")!=run_id
                    or receipt.get("input_sha256")!=input_sha256 or metadata.get("status") not in {"COMPLETED","COMPLETED_WITH_WARNINGS"}
                    or receipt.get("report_path")!=str(settings.report_path)):
                raise ReportUnavailable("Ce rapport n'est plus le rapport courant. Relancez l'analyse de votre archive pour le régénérer.")
            payload=settings.report_path.read_bytes()
            if not payload or hashlib.sha256(payload).hexdigest()!=receipt.get("report_sha256"):
                raise ReportUnavailable("Le rapport a été modifié ou est incomplet. Téléchargement bloqué.")
            return payload
        except (OSError,ValueError,KeyError) as exc:
            raise ReportUnavailable("Le rapport de cette exécution n'est pas disponible.") from exc
