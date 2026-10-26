"""Compare les sorties métier d'un notebook et d'une exécution App, sans GPU.

Les valeurs personnelles restent dans les répertoires comparés. Le diagnostic
affiche uniquement des chemins de divergence et compteurs, jamais les valeurs.
"""
import argparse
import hashlib
import json
from pathlib import Path
import sys

def read(path):return json.loads(path.read_text(encoding="utf-8"))

def field_projection(item):
    keys=("raw_value","normalized_value","status","usable","semantic_field","source_page",
          "source_evidence","detected_label","writing_type","readability","confidence",
          "model_confidence","manual_recheck_required","manual_recheck_reasons","extraction_method")
    return {key:item.get(key) for key in keys}

def documents(root):
    result={}
    for path in (root/"customers").glob("*/documents/*.json"):
        doc=read(path)
        key=(doc["customer_id"],doc["document_key"],doc["physical_document_id"])
        result[key]={
            "raw_pages":[{k:p.get(k) for k in ("page_number","raw_ocr","status","degenerate_output","stop_reason")}
                         for p in doc["raw_ocr"]["pages"]],
            "fields":{k:field_projection(v) for k,v in doc["fields"].items()},
            "forms":[{"form_type":f["form_type"],"page_number":f["page_number"],
                      "fields":{k:field_projection(v) for k,v in f["fields"].items()}} for f in doc["forms"]],
            "document_type":doc["document_type"],"document_found":doc["document_found"],
            "document_detection":doc.get("document_detection"),"validation":doc["validation"],
            "database_matching":doc.get("database_matching"),"manual_review":doc["manual_review"],
            "mrz":{k:doc.get("mrz",{}).get(k) for k in ("raw","present","format_valid","status")},
        }
    return result

def customers(root):
    result={}
    for path in (root/"customers").glob("*/*_extraction.json"):
        c=read(path)
        result[c["customer_id"]]={"reference_status":c["reference_status"],
            "database_consistency":c["database_consistency"],"manual_review":c["manual_review"],
            "cross_document_checks":c.get("cross_document_checks"),
            "family_statuses":{k:v["aggregate_status"] for k,v in c["document_families"].items()}}
    return result

def differences(left,right,path="",limit=100):
    result=[]
    def visit(a,b,p):
        if len(result)>=limit:return
        if type(a)!=type(b):result.append(p+":type");return
        if isinstance(a,dict):
            if a.keys()!=b.keys():result.append(p+":keys")
            for key in a.keys() & b.keys():visit(a[key],b[key],p+"/"+str(key))
        elif isinstance(a,list):
            if len(a)!=len(b):result.append(p+":length")
            for i,(x,y) in enumerate(zip(a,b)):visit(x,y,p+"/"+str(i))
        elif a!=b:result.append(p+":value")
    visit(left,right,path)
    # Hash customer/document keys before display; do not emit customer identities.
    return [hashlib.sha256(p.encode()).hexdigest()+":"+p.rsplit(":",1)[-1] for p in result]

def excel_business(path):
    from openpyxl import load_workbook
    w=load_workbook(path,read_only=True,data_only=False)
    result={"sheets":w.sheetnames,"business":{}}
    try:
        for name in ("CUSTOMER_CONTROLS","DOCUMENT_CONTROLS","FATCA_DETAILS","SIGNATURE_DETAILS",
                     "DATABASE_MATCHING","NON_CONFORMITIES","MANUAL_REVIEW_QUEUE"):
            s=w[name];rows=iter(s.values);header=next(rows)
            keep=[i for i,key in enumerate(header) if key and "PATH" not in str(key).upper()]
            result["business"][name]={"headers":[header[i] for i in keep],
                                     "rows":[[r[i] for i in keep] for r in rows]}
        return result
    finally:w.close()

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--notebook-output",type=Path,required=True)
    p.add_argument("--app-output",type=Path,required=True,help="output_qwen38/datasets/<SHA256>")
    p.add_argument("--notebook-excel",type=Path,required=True)
    p.add_argument("--app-excel",type=Path,required=True)
    args=p.parse_args()
    checks={"documents":(documents(args.notebook_output),documents(args.app_output)),
            "customers":(customers(args.notebook_output),customers(args.app_output)),
            "excel":(excel_business(args.notebook_excel),excel_business(args.app_excel))}
    failed=False
    for name,(left,right) in checks.items():
        diff=differences(left,right,name)
        if name in {"documents","customers"} and (not left or not right):diff.append("EMPTY_DATASET")
        print(name,"PASS" if not diff else "FAIL", "difference_paths_sha256",diff)
        if name in {"documents","customers"}:print("counts",len(left),len(right))
        failed |= bool(diff)
    return int(failed)

if __name__=="__main__":sys.exit(main())
